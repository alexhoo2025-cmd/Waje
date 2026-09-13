import datetime,hashlib,json,subprocess,sys
from pathlib import Path
from google.cloud import bigquery

P=Path(__file__).resolve().parent
ROOT=P.parents[1]
SOURCE=Path('/Users/robin/Desktop/waje data/提现大于充值9.10.xlsx')
SQL=P/'uidlog_aggregate_v2.sql'
sys.path.insert(0,str(ROOT/'analysis/paid_retention_aug_refresh_2026_09_09'))
from direct_client import client
sys.path.insert(0,str(ROOT/'scripts'))
from check_query_window import check

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()

reader='''
import json,sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],read_only=True,data_only=True);ws=wb['数据']
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))];idx={h:i for i,h in enumerate(headers)};out=[]
for row in ws.iter_rows(min_row=2,values_only=True):
 uid=row[idx['用户id']];tc=row[idx['时间段内TC差值']]
 if uid is not None and isinstance(tc,(int,float)):out.append({'uid':str(int(uid)),'tc_gap':float(tc),'recharge':float(row[idx['时间段内累计充值金额']]or 0),'withdraw':float(row[idx['时间段内累计TX金额']]or 0)})
print(json.dumps(out))
'''
bundled_python='/Users/robin/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3'
rows=json.loads(subprocess.check_output([bundled_python,'-c',reader,str(SOURCE)],text=True,timeout=120))
rows.sort(key=lambda x:x['tc_gap'],reverse=True)
top=rows[:30]
assert len(top)==30 and len({r['uid'] for r in top})==30
assert all(top[i]['tc_gap']>=top[i+1]['tc_gap'] for i in range(29))

window=check(SQL)
validator='/Users/robin/.codex/skills/waje-bigquery-readonly/scripts/validate_readonly_sql.py'
subprocess.run([sys.executable,validator,str(SQL)],check=True,capture_output=True,text=True)
params=[bigquery.ArrayQueryParameter('target_user_ids','STRING',[r['uid'] for r in top])]
config=bigquery.QueryJobConfig(query_parameters=params,use_query_cache=False,dry_run=True)
b=client()
dry=b.query(SQL.read_text(),job_config=config,location='europe-west4')
estimate=int(dry.total_bytes_processed or 0)
limit=5*1024**3
receipt={
 'status':'dry_run_complete',
 'source':{'path':str(SOURCE),'sha256':sha(SOURCE),'sheet':'数据','source_rows':len(rows)},
 'selection':{'method':'时间段内TC差值 descending, first 30','count':30,'cutoff_tc_gap':top[-1]['tc_gap'],'total_tc_gap':sum(r['tc_gap'] for r in top),'total_recharge':sum(r['recharge'] for r in top),'total_withdraw':sum(r['withdraw'] for r in top)},
 'query':{'sql':SQL.name,'sql_sha256':sha(SQL),'window_check':window,'date':'2026-09-10','app_id':90006,'event_type':'ASSET','user_ids_persisted':False,'estimated_bytes':estimate,'limit_bytes':limit},
 'created_at':datetime.datetime.now(datetime.timezone.utc).isoformat()
}
(P/'query-receipt-v2.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2,default=str))
print(json.dumps({'dry_run_gib':estimate/1024**3,'within_limit':estimate<=limit,'top30_tc_gap':receipt['selection']['total_tc_gap']},ensure_ascii=False),flush=True)
if estimate>limit:
    receipt['status']='blocked_scan_limit';(P/'query-receipt-v2.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2,default=str));raise SystemExit(2)

config=bigquery.QueryJobConfig(query_parameters=params,use_query_cache=False,maximum_bytes_billed=limit,labels={'purpose':'top30-uidlog-audit'})
job=b.query(SQL.read_text(),job_config=config,location='europe-west4')
result=[dict(r.items()) for r in job.result(timeout=180)]
assert len(result)<=3000 and all('user_id' not in r for r in result)
receipt['status']='completed'
receipt['query'].update(job_id=job.job_id,bytes_processed=job.total_bytes_processed,bytes_billed=job.total_bytes_billed,cache_hit=job.cache_hit,result_rows=len(result))
(P/'aggregate-result-v2.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
(P/'query-receipt-v2.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2,default=str))
print(json.dumps({'status':'completed','job_id':job.job_id,'bytes_processed':job.total_bytes_processed,'rows':len(result)},ensure_ascii=False))
