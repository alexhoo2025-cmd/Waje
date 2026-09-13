"""Read both research workbooks without modifying them; preserve cell lineage."""
from pathlib import Path
import openpyxl,json,hashlib,re,statistics
ROOT=Path(__file__).resolve().parent
FILES=[Path('/Users/robin/Desktop/坦桑调研/Research TZ _ Bet.xlsx'),Path('/Users/robin/Desktop/坦桑调研/坦桑行业调研数据.xlsx')]
def write(name,obj):
    (ROOT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str))
def main():
    inventory=[]; cells=[]; formulas=[]
    for fi,p in enumerate(FILES):
        w=openpyxl.load_workbook(p,read_only=True,data_only=False)
        values=openpyxl.load_workbook(p,read_only=True,data_only=True)
        cached={s.title:{c.coordinate:c.value for row in s for c in row if c.value is not None} for s in values}
        for s in w:
            n=0
            for row in s:
                for c in row:
                    if c.value is None: continue
                    n+=1
                    rec=dict(file=p.name,sheet=s.title,cell=c.coordinate,raw=c.value,format=c.number_format,kind=c.data_type,checked='2026-09-08')
                    if c.data_type=='f':
                        rec['cached']=cached[s.title].get(c.coordinate);formulas.append(rec)
                    cells.append(rec)
            inventory.append(dict(file=p.name,sheet=s.title,state=s.sheet_state,nonempty=n,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
        w.close();values.close()
    write('cells.json',cells);write('inventory.json',inventory);write('formulas.json',formulas)
    w=openpyxl.load_workbook(FILES[0],read_only=True,data_only=True)
    # Canonical identities are stable keys; suppliers never enter this list.
    names=['GSB','Premier Bet','Throne Bet','Meridian Bet','PM Bet','betPawa','SportyBet','SportPesa','Betway','Betika','Sokabet','WasafiBet','Gwala Bet','Mbet','WinPrincess Bet','888Bet','Parimatch','LeonBet','WiBet']
    records=[]
    for r,name in zip(range(4,23),names):
        row={'brand':name,'raw_brand':w['ONLINE'].cell(r,1).value,'source':f'ONLINE!A{r}:AB{r}'}
        for field,col in {'min_stake':3,'max_stake':4,'max_win':5,'min_deposit':6,'max_deposit':7,'welcome':8,'multibet':9,'cashout':10,'refund':11,'other_bonus':12,'dog':14,'horse':15,'virtual_football':16,'spin':18,'keno':19,'aviator':21,'casino':22,'jackpot_games':24,'jackpot_stake':25,'jackpot_amount':26,'ussd':28}.items():
            c=w['ONLINE'].cell(r,col);row[field]=c.value;row[field+'_cell']=c.coordinate
            if '%' in c.number_format and isinstance(c.value,(int,float)): row[field+'_display']=str(c.value*100)+'%'
        records.append(row)
    write('brands.json',records)
    def grid(sheet,maxr=100,maxc=30):
        s=w[sheet];return [{'row':r,'cells':{openpyxl.utils.get_column_letter(c):v for c,v in enumerate(row,1) if v is not None}} for r,row in enumerate(s.iter_rows(max_row=min(s.max_row,maxr),max_col=min(s.max_column,maxc),values_only=True),1) if any(v is not None for v in row)]
    for s in w.sheetnames: write('grid_'+s.strip().replace('/','_')+'.json',grid(s))
    stats={}
    for f in ['cashout','dog','horse','virtual_football','spin','keno','aviator','casino']:
        yes=sum(str(x[f]).strip().upper()=='YES' for x in records);no=sum(str(x[f]).strip().upper()=='NO' for x in records)
        stats[f]={'yes':yes,'no':no,'unknown_or_named':19-yes-no,'denominator':yes+no,'pct':round(100*yes/(yes+no),2) if yes+no else None}
    for f in ['min_stake','min_deposit']:
        vals=[x[f] for x in records if isinstance(x[f],(float,int))]
        stats[f]={'n':len(vals),'median':statistics.median(vals),'min':min(vals),'max':max(vals),'at_most100':sum(v<=100 for v in vals)}
    write('initial_stats.json',stats)
    print(json.dumps({'sheets':len(inventory),'cells':len(cells),'formulas':len(formulas),'brands':len(records),'stats':stats},ensure_ascii=False))
if __name__=='__main__':main()
