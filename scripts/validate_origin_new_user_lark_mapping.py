#!/usr/bin/env python3
"""Content-validate local Origin sheet aliases against a Lark backup."""

from __future__ import annotations

import argparse
import json
import math
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


MAPPINGS = [
    ("WajeSpecial-facebook", "WajeSpecial-facebook", "WajeSpecial-facebook.json"),
    ("WajeSpecial-googleadwords_int", "WajeSpecial-googleadwords_int", "WajeSpecial-googleadwords_int.json"),
    ("WajeSpecial-Google商店", "WajeSpecial-Google商店", "WajeSpecial-Google_.json"),
    ("wajeios-AppStore商店", "WAJEIOS-AppStore商店", "WAJEIOS-AppStore_.json"),
    ("wajebetH5-facebook", "WAJEBETH5", "WAJEBETH5.json"),
    ("wajeH5-fb", "wajeH5-facebook", "wajeH5-facebook.json"),
    ("wajeH5ga-googlewors_int", "wajeH5ga-googlewords_int", "wajeH5ga-googlewords_int.json"),
    ("pww", "PWA", "PWA.json"),
]


def iso(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value or "").strip().replace("/", "-")
    match = re.match(r"^(20\d\d)-(\d{1,2})-(\d{1,2})", text)
    return f"{match.group(1)}-{match.group(2).zfill(2)}-{match.group(3).zfill(2)}" if match else None


def header(value: object) -> str:
    return re.sub(r"[\s\u00a0\n\r，,()（）._-]+", "", str(value or "")).lower()


def value(value: object) -> object:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return text
    try:
        return float(text)
    except ValueError:
        return text


def same(a: object, b: object) -> bool:
    left, right = value(a), value(b)
    if left is None or right is None:
        return left is right
    if isinstance(left, float) and isinstance(right, float):
        return math.isclose(left, right, rel_tol=1e-8, abs_tol=1e-8)
    return left == right


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--local-workbook", required=True)
    parser.add_argument("--backup-cells", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sample-dates", default="2026-08-20,2026-08-21,2026-08-22,2026-08-23")
    args = parser.parse_args()
    sample_dates = args.sample_dates.split(",")
    workbook = load_workbook(args.local_workbook, data_only=True, read_only=False)
    output = {"schema_version": 1, "status": "ok", "sample_dates": sample_dates, "mappings": {}}
    for local_name, online_name, backup_file in MAPPINGS:
        worksheet = workbook[local_name]
        local_headers = [worksheet.cell(1, col).value for col in range(1, 44)]
        local_by_date = {iso(worksheet.cell(row, 1).value): [worksheet.cell(row, col).value for col in range(1, 44)] for row in range(2, worksheet.max_row + 1) if iso(worksheet.cell(row, 1).value)}
        payload = json.loads((Path(args.backup_cells) / backup_file).read_text())
        backup_cells = payload["ranges"][0]["cells"]
        online_headers = [(cell or {}).get("value") for cell in backup_cells[0][:43]]
        online_by_date = {iso((row[0] or {}).get("value")): [(cell or {}).get("value") for cell in row[:43]] for row in backup_cells[1:] if row and iso((row[0] or {}).get("value"))}
        header_diffs = [{"column_index": i + 1, "local": local_headers[i], "online": online_headers[i]} for i in range(43) if header(local_headers[i]) != header(online_headers[i])]
        samples = []
        for day in sample_dates:
            local_row = local_by_date.get(day)
            online_row = online_by_date.get(day)
            if local_row is None or online_row is None:
                samples.append({"date": day, "status": "missing"})
                continue
            diffs = [{"column_index": i + 1, "local": local_row[i], "online": online_row[i]} for i in range(43) if not ((iso(local_row[i]) == iso(online_row[i])) if i == 0 else same(local_row[i], online_row[i]))]
            samples.append({"date": day, "status": "match" if not diffs else "mismatch", "mismatch_count": len(diffs), "mismatches": diffs})
        content_ok = all(item["status"] == "match" for item in samples)
        allowed_header_alias = local_name == "pww" and len(header_diffs) == 1 and header_diffs[0]["column_index"] == 1
        header_ok = not header_diffs or allowed_header_alias
        status = "validated" if content_ok and header_ok else "blocked"
        output["mappings"][local_name] = {"online_sheet": online_name, "backup_file": backup_file, "status": status, "online_column_count": len(backup_cells[0]), "header_differences": header_diffs, "header_alias_accepted_by_content": allowed_header_alias, "sample_comparisons": samples, "extra_columns_preserved": len(backup_cells[0]) > 43}
        if status != "validated":
            output["status"] = "blocked"
    Path(args.output).write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps({"status": output["status"], "validated": sum(x["status"] == "validated" for x in output["mappings"].values()), "total": len(MAPPINGS)}, ensure_ascii=False))
    if output["status"] != "ok":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
