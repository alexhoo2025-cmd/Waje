#!/usr/bin/env python3
"""Archive and validate one Origin BQ-新增付费用户分析 export.

The Origin XLSX files occasionally contain a stale worksheet dimension.  This
reader deliberately uses normal (non-streaming) workbook mode and validates
the actual 43-column matrix instead of trusting that dimension metadata.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


HEADERS = [
    "日期", "区服", "新增人数", "终身", "首日", "次日", "3日", "4日", "5日", "6日", "7日", "8日", "9日", "10日", "11日", "12日", "13日", "14日", "15日", "30日", "60日",
    "新增付费率", "新增付费人数", "次留", "3日留", "7日留", "15日留", "30日留", "60日留", "tc比", "tx率", "人均tx金额", "首充付费率", "首充付费人数", "首充次留", "首充3日留", "首充7日留", "首充15日留", "首充30日留", "首充60日留", "首充tc比", "首充tx率", "首充人均tx金额",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_header(value: object) -> str:
    return "" if value is None else "".join(str(value).split())


def iso_date(value: object) -> str | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], pattern).date().isoformat()
        except ValueError:
            pass
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--raw-root", required=True)
    parser.add_argument("--sheet-key", required=True)
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    parser.add_argument("--package", required=True)
    parser.add_argument("--media", default="")
    parser.add_argument("--channel", default="")
    parser.add_argument("--receipt-json", default="")
    args = parser.parse_args()

    source = Path(args.input).expanduser().resolve()
    if not source.is_file():
        raise SystemExit(f"input does not exist: {source}")
    destination_dir = Path(args.raw_root).expanduser().resolve() / args.sheet_key
    destination = destination_dir / "source.xlsx"
    snapshot = destination_dir / "source-data.json"
    receipt = destination_dir / "query-receipt.json"
    if destination.exists() or snapshot.exists() or receipt.exists():
        raise SystemExit(f"refusing to overwrite immutable raw snapshot: {destination_dir}")

    workbook = load_workbook(source, data_only=True, read_only=False)
    worksheet = workbook.active
    headers = [worksheet.cell(1, index + 1).value for index in range(len(HEADERS))]
    if [normalize_header(value) for value in headers] != [normalize_header(value) for value in HEADERS]:
        raise SystemExit("43 Origin headers do not match the accepted report contract")

    rows = []
    for row_number in range(2, worksheet.max_row + 1):
        row = [worksheet.cell(row_number, column + 1).value for column in range(len(HEADERS))]
        date = iso_date(row[0])
        if not date:
            continue
        row[0] = date
        rows.append(row)
    dates = [row[0] for row in rows]
    expected = []
    current = datetime.strptime(args.start_date, "%Y-%m-%d").date()
    end = datetime.strptime(args.end_date, "%Y-%m-%d").date()
    from datetime import timedelta
    while current <= end:
        expected.append(current.isoformat())
        current += timedelta(days=1)
    if sorted(dates) != expected or len(set(dates)) != len(dates):
        raise SystemExit(f"date coverage invalid: expected {expected[0]}..{expected[-1]} ({len(expected)}); got {sorted(dates)}")
    if any(len(row) != len(HEADERS) for row in rows):
        raise SystemExit("one or more source rows are not 43 columns")

    destination_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    payload = {
        "schema_version": 1,
        "source": "Origin BQ-新增付费用户分析",
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "origin_export": {
            "input_path": str(source),
            "sha256": sha256(source),
            "archived_path": str(destination),
            "archived_sha256": sha256(destination),
        },
        "query": {
            "start_date": args.start_date,
            "end_date": args.end_date,
            "package": args.package,
            "media": args.media or None,
            "channel": args.channel or None,
            "tc_logic": "累计利润(C-T)",
        },
        "headers": headers,
        "rows": rows,
        "validation": {
            "status": "ok",
            "date_count": len(dates),
            "date_range": {"start": min(dates), "end": max(dates)},
            "field_count": len(headers),
            "data_rows": len(rows),
        },
    }
    receipt_payload = {
        "schema_version": 1,
        "status": "source_export_validated",
        "sheet_key": args.sheet_key,
        "query": payload["query"],
        "export_sha256": payload["origin_export"]["sha256"],
        "observed_date_count": len(dates),
        "observed_field_count": len(headers),
        "sample": {"newest": rows[0][:7], "oldest": rows[-1][:7]},
        "note": "Raw XLSX kept intact; JSON is a validated logical snapshot derived from its 43 fields.",
    }
    if args.receipt_json:
        extra = json.loads(Path(args.receipt_json).read_text())
        receipt_payload["ui_receipt"] = extra
    snapshot.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    receipt.write_text(json.dumps(receipt_payload, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"status": "ok", "sheet_key": args.sheet_key, "source_sha256": payload["origin_export"]["sha256"], "rows": len(rows), "first": rows[0][:7], "last": rows[-1][:7]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
