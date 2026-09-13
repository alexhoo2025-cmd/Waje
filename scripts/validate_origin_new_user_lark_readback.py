#!/usr/bin/env python3
"""Verify an Origin new-user Lark publish against its backed-up source workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


MAPPINGS = [
    ("WajeSpecial-facebook", "WajeSpecial-facebook", "WajeSpecial-facebook.json"),
    ("WajeSpecial-googleadwords_int", "WajeSpecial-googleadwords_int", "WajeSpecial-googleadwords_int.json"),
    ("WajeSpecial-Google商店", "WajeSpecial-Google商店", "WajeSpecial-Google_.json"),
    ("wajeios-AppStore商店", "WAJEIOS-AppStore商店", "WAJEIOS-AppStore_.json"),
    ("wajebetH5-facebook", "WAJEBETH5", "WAJEBETH5.json"),
    ("wajeH5-fb", "wajeH5-facebook", "wajeH5-facebook.json"),
    ("wajeH5ga-googlewors_int", "wajeH5ga-googlewords_int", "wajeH5ga-googlewords_int.json"),
    ("pww", "PWA", "PWA.json"),
]

ERROR_RE = re.compile(r"#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A", re.I)


def iso(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)) and 20000 < value < 100000:
        # Excel 1900 date system, sufficient for this 2026 workbook range.
        return (datetime(1899, 12, 30) + timedelta(days=round(value))).date().isoformat()
    text = str(value or "").strip().replace("/", "-")
    match = re.match(r"^(20\d\d)-(\d{1,2})-(\d{1,2})", text)
    return f"{match.group(1)}-{match.group(2).zfill(2)}-{match.group(3).zfill(2)}" if match else None


def scalar(value: object) -> object:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return text
    try:
        return float(text)
    except ValueError:
        return text


def same(left: object, right: object) -> bool:
    a, b = scalar(left), scalar(right)
    if a is None or b is None:
        return a is b
    if isinstance(a, float) and isinstance(b, float):
        return math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-8)
    return a == b


def same_with_display_precision(local: object, remote_cell: object) -> bool:
    """Lark readback returns formatted display values, not binary cell values."""
    remote = cell_value(remote_cell)
    a, b = scalar(local), scalar(remote)
    if a is None or b is None:
        return a is b
    if not (isinstance(a, float) and isinstance(b, float)):
        return a == b
    number_format = str(((remote_cell or {}).get("cell_styles") or {}).get("number_format") or "")
    if "%" in number_format or str(remote).strip().endswith("%"):
        return math.isclose(a, b, rel_tol=1e-8, abs_tol=0.000051)  # 0.005 percentage point plus binary rounding
    if "0.##" in number_format or "0.00" in number_format:
        return math.isclose(a, b, rel_tol=1e-8, abs_tol=0.005)
    return math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-8)


def cell_value(cell: object) -> object:
    return (cell or {}).get("value") if isinstance(cell, dict) else None


def rows_from_backup(payload: dict) -> list[dict]:
    output = []
    for index, cells in enumerate(payload["ranges"][0]["cells"][1:], 2):
        day = iso(cell_value(cells[0]) if cells else None)
        if day:
            output.append({"row": index, "date": day, "cells": cells})
    return output


def values(cells: list[dict]) -> list[object]:
    return [cell_value(cell) for cell in cells]


def style(cell: object) -> dict:
    cell = cell or {}
    return {"cell_styles": cell.get("cell_styles"), "border_styles": cell.get("border_styles")}


def stable_hash(value: object) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def layout_shape(payload: dict) -> dict:
    data = payload.get("data", {})
    return {key: data.get(key) for key in ["column_groups", "column_widths", "frozen_columns", "frozen_rows", "hidden_columns", "hidden_rows", "merged_cells", "row_groups"]}


def conditional_shape(payload: dict) -> dict:
    items = payload.get("data", {}).get("sheets", [{}])[0].get("conditional_formats", [])
    return {item.get("conditional_format_id"): {"rule_type": item.get("details", {}).get("rule_type"), "attrs": item.get("details", {}).get("attrs"), "style": item.get("details", {}).get("style")} for item in items}


def row_height_at(layout: dict, row: int) -> tuple[object, object] | None:
    for item in layout.get("data", {}).get("row_heights", []):
        match = re.match(r"^(\d+):(\d+)$", str(item.get("rows", "")))
        if match and int(match.group(1)) <= row <= int(match.group(2)):
            return item.get("height"), item.get("type")
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--local-workbook", required=True)
    parser.add_argument("--before-root", required=True)
    parser.add_argument("--after-root", required=True)
    parser.add_argument("--run-dir", required=True)
    parser.add_argument("--history-cutoff", required=True)
    args = parser.parse_args()

    run_dir = Path(args.run_dir)
    receipt = json.loads((run_dir / "run-receipt.json").read_text())
    accepted = receipt["accepted_dates"]
    excluded = receipt["excluded_not_mature_dates"]
    before_root = Path(args.before_root)
    after_root = Path(args.after_root)
    before_manifest = json.loads((before_root / "backup-manifest.json").read_text())
    after_manifest = json.loads((after_root / "backup-manifest.json").read_text())
    if not before_manifest["integrity"]["complete"] or not after_manifest["integrity"]["complete"]:
        raise SystemExit("incomplete Lark backup prevents verification")
    workbook = load_workbook(args.local_workbook, data_only=True, read_only=False)

    output: dict[str, object] = {
        "schema_version": 1,
        "status": "ok",
        "revision_before": before_manifest["revision"],
        "revision_after": after_manifest["revision"],
        "accepted_dates": accepted,
        "excluded_not_mature_dates": excluded,
        "history_cutoff_exclusive": args.history_cutoff,
        "sheets": {},
    }
    history_report: dict[str, object] = {"status": "ok", "history_cutoff_exclusive": args.history_cutoff, "sheets": {}}
    format_report: dict[str, object] = {"status": "ok", "sheets": {}}
    formula_errors = []
    readback = {"schema_version": 1, "status": "ok", "sheets": {}}

    for local_name, online_name, backup_file in MAPPINGS:
        before = json.loads((before_root / "cells" / backup_file).read_text())
        after = json.loads((after_root / "cells" / backup_file).read_text())
        before_rows = rows_from_backup(before)
        after_rows = rows_from_backup(after)
        before_by_date = {item["date"]: item for item in before_rows}
        after_by_date: dict[str, list[dict]] = {}
        for item in after_rows:
            after_by_date.setdefault(item["date"], []).append(item)
        worksheet = workbook[local_name]
        local_by_date = {iso(worksheet.cell(row, 1).value): [worksheet.cell(row, col).value for col in range(1, 44)] for row in range(2, worksheet.max_row + 1) if iso(worksheet.cell(row, 1).value)}
        problems = []
        targets = []
        for day in accepted:
            online_rows = after_by_date.get(day, [])
            if len(online_rows) != 1:
                problems.append({"date": day, "error": "date_missing_or_duplicate", "count": len(online_rows)})
                continue
            local = local_by_date.get(day)
            if local is None:
                problems.append({"date": day, "error": "missing_local_output"})
                continue
            remote_cells = online_rows[0]["cells"][:43]
            remote = values(remote_cells)
            diffs = []
            for index in range(43):
                correct = iso(local[index]) == iso(remote[index]) if index == 0 else same_with_display_precision(local[index], remote_cells[index])
                if not correct:
                    diffs.append({"column_index": index + 1, "local": local[index], "online": remote[index]})
            if diffs:
                problems.append({"date": day, "error": "value_mismatch", "mismatches": diffs})
            targets.append({"date": day, "row": online_rows[0]["row"], "value_hash": stable_hash(remote)})
        excluded_present = [day for day in excluded if day in after_by_date]
        if excluded_present:
            problems.append({"error": "unmatured_dates_present", "dates": excluded_present})

        history_before = {item["date"]: values(item["cells"]) for item in before_rows if item["date"] < args.history_cutoff}
        history_after = {items[0]["date"]: values(items[0]["cells"]) for items in after_by_date.values() if len(items) == 1 and items[0]["date"] < args.history_cutoff}
        history_same = stable_hash(history_before) == stable_hash(history_after)
        if not history_same:
            problems.append({"error": "historical_prefix_changed"})
            history_report["status"] = "blocked"
        history_report["sheets"][online_name] = {"unchanged": history_same, "before_hash": stable_hash(history_before), "after_hash": stable_hash(history_after)}

        # Existing dates retain all original cell styles.  Appended rows inherit
        # the last pre-existing row style across all original columns.
        existing_style_same = True
        for day in accepted:
            if day in before_by_date and len(after_by_date.get(day, [])) == 1:
                left = before_by_date[day]["cells"]
                right = after_by_date[day][0]["cells"]
                if [style(c) for c in left] != [style(c) for c in right]:
                    existing_style_same = False
                    break
        last_before = max(before_rows, key=lambda item: item["row"])
        appended_style_same = True
        for day in accepted:
            if day not in before_by_date and len(after_by_date.get(day, [])) == 1:
                if [style(c) for c in after_by_date[day][0]["cells"]] != [style(c) for c in last_before["cells"]]:
                    appended_style_same = False
                    break
        before_layout = json.loads((before_root / "layout" / backup_file).read_text())
        after_layout = json.loads((after_root / "layout" / backup_file).read_text())
        layout_same = layout_shape(before_layout) == layout_shape(after_layout)
        existing_row_heights_same = all(
            row_height_at(before_layout, before_by_date[day]["row"]) == row_height_at(after_layout, after_by_date[day][0]["row"])
            for day in accepted if day in before_by_date and len(after_by_date.get(day, [])) == 1
        )
        appended_row_heights_same = all(
            row_height_at(before_layout, last_before["row"]) == row_height_at(after_layout, after_by_date[day][0]["row"])
            for day in accepted if day not in before_by_date and len(after_by_date.get(day, [])) == 1
        )
        before_cond = json.loads((before_root / "conditional-format" / backup_file).read_text())
        after_cond = json.loads((after_root / "conditional-format" / backup_file).read_text())
        cond_same = conditional_shape(before_cond) == conditional_shape(after_cond)
        if not (existing_style_same and appended_style_same and existing_row_heights_same and appended_row_heights_same and layout_same and cond_same):
            format_report["status"] = "blocked"
            problems.append({"error": "format_or_layout_changed", "existing_style_same": existing_style_same, "appended_style_same": appended_style_same, "existing_row_heights_same": existing_row_heights_same, "appended_row_heights_same": appended_row_heights_same, "layout_same": layout_same, "conditional_rules_same": cond_same})
        format_report["sheets"][online_name] = {"existing_style_same": existing_style_same, "appended_style_same": appended_style_same, "existing_row_heights_same": existing_row_heights_same, "appended_row_heights_same": appended_row_heights_same, "layout_same": layout_same, "conditional_rules_same": cond_same, "before_row_count": len(before_rows) + 1, "after_row_count": len(after_rows) + 1}

        for row in after_rows:
            for col, cell in enumerate(row["cells"], 1):
                candidate = str(cell_value(cell) or "")
                formula = str((cell or {}).get("formula") or "")
                if ERROR_RE.search(candidate) or ERROR_RE.search(formula):
                    formula_errors.append({"sheet": online_name, "cell": f"r{row['row']}c{col}", "value": candidate, "formula": formula})
        if problems:
            output["status"] = "blocked"
        output["sheets"][online_name] = {"local_sheet": local_name, "status": "ok" if not problems else "blocked", "accepted_date_rows": targets, "problems": problems, "extra_columns_preserved_on_existing_rows": all(values(before_by_date[day]["cells"])[43:] == values(after_by_date[day][0]["cells"])[43:] for day in accepted if day in before_by_date and len(after_by_date.get(day, [])) == 1)}
        readback["sheets"][online_name] = {"accepted_dates": targets, "date_count": len(after_rows), "last_date": max(after_by_date) if after_by_date else None}

    if formula_errors:
        output["status"] = "blocked"
    formula_report = {"status": "ok" if not formula_errors else "blocked", "error_count": len(formula_errors), "errors": formula_errors}
    if history_report["status"] != "ok" or format_report["status"] != "ok":
        output["status"] = "blocked"
    readback["status"] = output["status"]
    (run_dir / "readback-snapshot.json").write_text(json.dumps(readback, ensure_ascii=False, indent=2, default=str) + "\n")
    (run_dir / "historical-integrity-report.json").write_text(json.dumps(history_report, ensure_ascii=False, indent=2, default=str) + "\n")
    (run_dir / "format-integrity-report.json").write_text(json.dumps(format_report, ensure_ascii=False, indent=2, default=str) + "\n")
    (run_dir / "formula-verification.json").write_text(json.dumps(formula_report, ensure_ascii=False, indent=2, default=str) + "\n")
    (run_dir / "lark-readback-validation.json").write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps({"status": output["status"], "formula_errors": len(formula_errors), "history": history_report["status"], "format": format_report["status"]}, ensure_ascii=False))
    if output["status"] != "ok":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
