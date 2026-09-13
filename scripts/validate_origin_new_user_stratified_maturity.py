#!/usr/bin/env python3
"""Validate row-level Origin publication with field-level cohort maturity."""

from __future__ import annotations

import argparse
import json
import math
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


SHEETS = [
    "WajeSpecial-facebook", "WajeSpecial-googleadwords_int", "WajeSpecial-Google商店",
    "wajeios-AppStore商店", "wajebetH5-facebook", "pww", "wajeH5-fb", "wajeH5ga-googlewors_int",
]
WINDOWS = {
    5: 1, 6: 3, 7: 4, 8: 5, 9: 6, 10: 7, 11: 8, 12: 9, 13: 10, 14: 11,
    15: 12, 16: 13, 17: 14, 18: 15, 19: 30, 20: 60,
    23: 1, 24: 3, 25: 7, 26: 15, 27: 30, 28: 60,
    34: 1, 35: 3, 36: 7, 37: 15, 38: 30, 39: 60,
}


def iso(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value or "").strip().replace("/", "-")
    parts = text.split("-")
    if len(parts) >= 3 and len(parts[0]) == 4 and parts[0].startswith("20"):
        return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2][:2].zfill(2)}"
    return None


def scalar(value: object) -> object:
    if value is None or str(value).strip() in {"", "-"}:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return text
    try:
        return float(text)
    except ValueError:
        return text


def same(left: object, right: object) -> bool:
    a, b = scalar(left), scalar(right)
    if a is None or b is None:
        return a is b
    if isinstance(a, float) and isinstance(b, float):
        return math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-8)
    return a == b


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--raw-dir", required=True)
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    parser.add_argument("--maturity-cutoff", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    start, end, cutoff = (date.fromisoformat(value) for value in [args.start_date, args.end_date, args.maturity_cutoff])
    dates = []
    current = start
    while current <= end:
        dates.append(current.isoformat())
        current += timedelta(days=1)
    workbook = load_workbook(args.workbook, data_only=True, read_only=False)
    report = {"schema_version": 1, "status": "ok", "row_policy": "stratified", "dates": dates, "maturity_cutoff": args.maturity_cutoff, "sheets": {}}
    for sheet_name in SHEETS:
        raw = json.loads((Path(args.raw_dir) / f"{sheet_name}.json").read_text())
        raw_by_date = {iso(row[0]): row for row in raw["rows"]}
        ws = workbook[sheet_name]
        local_by_date = {iso(ws.cell(row, 1).value): [ws.cell(row, column).value for column in range(1, 44)] for row in range(2, ws.max_row + 1) if iso(ws.cell(row, 1).value)}
        failures, pending, mature_zeroes = [], 0, 0
        for day in dates:
            source = raw_by_date.get(day)
            output = local_by_date.get(day)
            if source is None or output is None:
                failures.append({"date": day, "error": "missing_source_or_output"})
                continue
            for index in range(43):
                window = WINDOWS.get(index)
                should_blank = window is not None and date.fromisoformat(day) + timedelta(days=window) > cutoff
                actual = output[index]
                if should_blank:
                    pending += 1
                    if scalar(actual) is not None:
                        failures.append({"date": day, "column_index": index + 1, "error": "unmatured_value_present", "actual": actual})
                else:
                    matches = iso(actual) == iso(source[index]) if index == 0 else same(actual, source[index])
                    if not matches:
                        failures.append({"date": day, "column_index": index + 1, "error": "mature_or_core_value_mismatch", "source": source[index], "actual": actual})
                    if scalar(source[index]) == 0:
                        mature_zeroes += 1
                        if scalar(actual) != 0:
                            failures.append({"date": day, "column_index": index + 1, "error": "mature_zero_not_preserved", "source": source[index], "actual": actual})
        report["sheets"][sheet_name] = {"status": "ok" if not failures else "blocked", "pending_window_cells": pending, "mature_zeroes_checked": mature_zeroes, "failures": failures}
        if failures:
            report["status"] = "blocked"
    Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps({"status": report["status"], "sheet_count": len(SHEETS), "pending_window_cells": sum(item["pending_window_cells"] for item in report["sheets"].values())}, ensure_ascii=False))
    if report["status"] != "ok":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
